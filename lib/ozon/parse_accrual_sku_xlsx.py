#!/usr/bin/env python3
"""Durable parse of Ozon «Отчёт по начислениям» -> JSON facts (company-bound ingest input)."""
import json, sys, hashlib, os
from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP

try:
    import openpyxl
except ImportError:
    os.system(f"{sys.executable} -m pip install -q openpyxl")
    import openpyxl

def D(x):
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))

def parse_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    s = str(v)[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").date().isoformat()
    except Exception:
        return ""

def classify(group, typ):
    g = (group or "").strip().lower().replace("ё", "е")
    t = (typ or "").strip().lower().replace("ё", "е")
    if "продаж" in g:
        if "балл" in t: return "points_sale"
        if "партн" in t: return "partner_sale"
        if "выруч" in t: return "revenue_sale"
        return "sales_other"
    if "возврат" in g:
        if "балл" in t: return "points_return"
        if "партн" in t: return "partner_return"
        if "выруч" in t: return "revenue_return"
        if "вознагражд" in t: return "commission_return"
        return "return_other"
    if "вознаграждение" in g:
        if "возврат" in t: return "commission_return"
        return "commission"
    if "услуги доставки" in g or "доставк" in g:
        return "logistics"
    if "реклам" in g or "продвиж" in g:
        return "ads"
    if "fbo" in g or "фбо" in g:
        return "fbo"
    if "услуг" in g or "эквайр" in t or "звезд" in t or "drop" in t or "упаков" in t or "обеспеч" in t or "обработк" in t:
        return "other_sku"
    return "other"

def parse_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    agg = defaultdict(lambda: defaultdict(lambda: Decimal("0")))
    n = 0
    no_sku = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        vals = list(row)
        if not vals or vals[0] is None:
            continue
        d = parse_date(vals[1])
        if not d:
            continue
        article = str(vals[4] or "").strip()
        sku = str(vals[5] or "").strip()
        group = str(vals[2] or "")
        typ = str(vals[3] or "")
        amt = Decimal(str(vals[15] or 0))
        cat = classify(group, typ)
        n += 1
        if not article and not sku:
            no_sku += 1
            continue
        key = (d, article, sku)
        agg[key][cat] += amt
        agg[key]["_n"] += 1
    wb.close()
    facts = []
    for (d, vendor, sku), c in agg.items():
        rev = c["revenue_sale"] + c["revenue_return"]
        points = c["points_sale"] + c["points_return"]
        partner = c["partner_sale"] + c["partner_return"]
        comm = -(c["commission"] + c["commission_return"])
        log = -c["logistics"]
        other = -(c["other_sku"] + c["return_other"] + c["sales_other"] + c["other"])
        fbo = -c["fbo"]
        ads = -c["ads"]
        facts.append({
            "operationDate": d,
            "vendorCode": vendor,
            "sku": sku,
            "taxableRevenue": D(rev),
            "discountPoints": D(points),
            "partnerPrograms": D(partner),
            "commissionExpense": D(comm),
            "logisticsExpense": D(log),
            "otherSkuExpense": D(other),
            "fboExpense": D(fbo),
            "adsExpense": D(ads),
            "rowsCount": int(c["_n"]),
        })
    return {"parsedRows": n, "skippedNoSku": no_sku, "facts": facts}

def main():
    path = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "/tmp/accrual_sku_facts_parsed.json"
    data = parse_file(path)
    data["sourceFileName"] = os.path.basename(path)
    with open(path, "rb") as f:
        data["sourceFileHash"] = hashlib.sha256(f.read()).digest().hex()
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    print(json.dumps({"out": out, "facts": len(data["facts"]), "parsedRows": data["parsedRows"], "skippedNoSku": data["skippedNoSku"]}))

if __name__ == "__main__":
    main()
